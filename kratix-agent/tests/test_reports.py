"""B2: report files (Excel and PowerPoint) from ERP data, the create_report tool, and the report routes."""

import asyncio
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from pptx import Presentation

from app.auth import CurrentUser
from app.erp_client import ErpError
from app.reports.catalog import REPORT_TYPES
from app.reports.service import ReportError, ReportService
from app.reports.slides import ROWS_PER_SLIDE
from app.store import Store
from app.tools import ToolContext, run_tool
from tests.test_tools import FakeErp

OWNER = CurrentUser(1, "owner", "Owner", "token-owner")
SALES = CurrentUser(2, "sana", "Sales", "token-sales")
WORKER = CurrentUser(3, "worker", "Worker", "token-worker")
NOW = datetime(2026, 10, 5, 9, 0, tzinfo=UTC)


def invoices(n):
    return [{"invoiceNumber": f"INV-{i:03d}", "invoiceDate": "2026-10-03T00:00:00", "customerName": f"Customer {i % 3}",
             "status": "Paid", "netAmount": 1000.0 * i, "paid": 1000.0 * i, "balance": 0.0} for i in range(1, n + 1)]


ERP_DATA = {
    "reports/sales": {
        "from": "2026-10-01T00:00:00", "to": "2026-10-05T00:00:00",
        "totals": {"invoices": 2, "subtotal": 318840.0, "tax": 0.0, "net": 318840.0, "paid": 318840.0, "balance": 0.0},
        "byCustomer": [{"customerName": "Karachi Wholesale Mart", "invoices": 1, "net": 311000.0, "paid": 311000.0, "balance": 0.0},
                       {"customerName": "Ali Sons Retail", "invoices": 1, "net": 7840.0, "paid": 7840.0, "balance": 0.0}],
        "byProduct": [{"productName": "Chakki Atta 10kg", "unit": "bag", "quantity": 500.0, "amount": 275000.0},
                      {"productName": "Suji 20kg", "unit": "bag", "quantity": 56.0, "amount": 7840.0}],
        "invoices": invoices(20),
    },
    "reports/receivables": {
        "asOf": "2026-10-05T00:00:00+05:00",
        "totals": {"notDue": 100.0, "days1To30": 200.0, "days31To60": 0.0, "days61Plus": 50.0, "total": 350.0},
        "byCustomer": [{"customerName": "Ali Sons Retail", "invoices": 2,
                        "aging": {"notDue": 100.0, "days1To30": 200.0, "days31To60": 0.0, "days61Plus": 50.0, "total": 350.0}}],
        "invoices": [{"invoiceNumber": "INV-9", "invoiceDate": "2026-07-01T00:00:00", "dueDate": "2026-07-31T00:00:00",
                      "customerName": "Ali Sons Retail", "netAmount": 50.0, "balance": 50.0, "daysOverdue": 66,
                      "bucket": "Days61Plus"}],
    },
    "reports/stock": {"from": "2026-10-01T00:00:00", "to": "2026-10-05T00:00:00", "rows": [
        {"productCode": "ATTA-10", "productName": "Chakki Atta 10kg", "category": "Finished Good", "stockForm": "Standard",
         "unit": "bag", "opening": 0.0, "received": 0.0, "produced": 1200.0, "used": 0.0, "dispatched": 500.0, "closing": 700.0},
        {"productCode": "WHT-A", "productName": "Wheat Grade A", "category": "Raw Material", "stockForm": "Standard",
         "unit": "kg", "opening": 0.0, "received": 20000.0, "produced": 0.0, "used": 20000.0, "dispatched": 0.0, "closing": 0.0},
    ]},
    "inventory": [
        {"productName": "Chakki Atta 10kg", "category": "Finished Good", "unit": "bag", "warehouseName": "FG Warehouse",
         "lotNumber": "MB-2026-001", "quantityAvailable": 700.0},
        {"productName": "Wheat Grade A", "category": "Raw Material", "unit": "kg", "warehouseName": "Godown A",
         "lotNumber": "LOT-1", "quantityAvailable": 0.0},
    ],
    "reports/production": {
        "from": "2026-10-01T00:00:00", "to": "2026-10-05T00:00:00",
        "totals": {"batches": 2, "input": 40000.0, "output": 38400.0, "waste": 1600.0, "yield": 96.0},
        "daily": [{"date": "2026-10-03T00:00:00", "batches": 1, "input": 20000.0, "output": 19200.0, "waste": 800.0, "yield": 96.0},
                  {"date": "2026-10-04T00:00:00", "batches": 1, "input": 20000.0, "output": 19200.0, "waste": 800.0, "yield": 96.0}],
        "batches": [{"batchID": 1, "batchNumber": "MB-2026-001", "productionDate": "2026-10-03T00:00:00", "shiftNumber": 1,
                     "operatorName": "Nadeem Abbas", "input": 20000.0, "output": 19200.0, "waste": 800.0, "yield": 96.0}],
        "outputs": [{"productName": "Chakki Atta 10kg", "materialType": "Output", "quantityKg": 24000.0}],
    },
}


@pytest.fixture
def service(tmp_path):
    clock = [NOW]
    svc = ReportService(Store(":memory:"), tmp_path, "Crescent Flour Mills", retention_days=7, per_hour=10,
                        clock=lambda: clock[0])
    svc.clock = clock
    svc.root = tmp_path
    return svc


def make(service, user=OWNER, kind="sales", fmt="xlsx", erp=None, **dates):
    return asyncio.run(service.create(user, erp or FakeErp(ERP_DATA), kind, fmt, **dates))


def path_of(service, report):
    return service.root / "reports" / str(report["userId"]) / f"{report['id']}.{report['format']}"


@pytest.mark.parametrize("kind", list(REPORT_TYPES))
@pytest.mark.parametrize("fmt", ["xlsx", "pptx"])
def test_every_report_type_opens_in_both_formats(service, kind, fmt):
    report = make(service, kind=kind, fmt=fmt, period="this_month")
    path = path_of(service, report)
    assert path.is_file() and report["sizeBytes"] == path.stat().st_size
    if fmt == "xlsx":
        assert load_workbook(path).sheetnames[0] == "Summary"
    else:
        assert len(Presentation(path).slides) >= 3


def test_sales_excel_totals_match_the_erp(service):
    report = make(service, period="this_month")
    book = load_workbook(path_of(service, report))
    summary = {row[0].value: row[1].value for row in book["Summary"].iter_rows(min_row=5) if row[0].value}
    assert summary["Net sales"] == 318840.0 and summary["Invoices"] == 2
    by_customer = list(book["By customer"].iter_rows(min_row=2, values_only=True))
    assert by_customer[0][:3] == ("Karachi Wholesale Mart", 1, 311000.0)
    assert book["By customer"].freeze_panes == "A2"
    assert book["Invoices"].max_row == 21
    assert book["Summary"]["A1"].value == "Crescent Flour Mills - Sales Report"
    assert report["fileName"] == "Sales_2026-10-01_to_2026-10-05.xlsx"
    assert report["params"] == {"from": "2026-10-01", "to": "2026-10-05"}


def test_receivables_is_as_of_today(service):
    report = make(service, kind="receivables", period="last_month")
    assert report["fileName"] == "Receivables_as_of_2026-10-05.xlsx"
    rows = list(load_workbook(path_of(service, report))["Open invoices"].iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "INV-9" and rows[0][-1] == "61+ days"


def test_deck_has_native_charts_and_continues_long_tables(service):
    deck = Presentation(path_of(service, make(service, fmt="pptx")))
    titles = [s.shapes[0].text_frame.text for s in list(deck.slides)[1:]]
    assert "Sales by customer (PKR)" in titles and "Sales by product (PKR)" in titles
    assert sum(shape.has_chart for s in deck.slides for shape in s.shapes) == 2
    assert "Invoices (1 of 2)" in titles and "Invoices (2 of 2)" in titles
    table = next(shape.table for s in deck.slides for shape in s.shapes
                 if shape.has_table and s.shapes[0].text_frame.text == "Invoices (1 of 2)")
    assert len(table.rows) == ROWS_PER_SLIDE + 1


def test_stock_charts_keep_bags_and_kg_apart(service):
    deck = Presentation(path_of(service, make(service, kind="stock", fmt="pptx")))
    titles = [s.shapes[0].text_frame.text for s in deck.slides]
    assert "Closing stock (bags)" in titles and "Closing stock (kg)" not in titles  # wheat closing is 0


def test_explicit_dates_and_bad_input(service):
    report = make(service, kind="production", from_date="2026-09-15", to_date="2026-09-30")
    assert report["params"] == {"from": "2026-09-15", "to": "2026-09-30"}
    for kwargs, message in [({"kind": "payroll"}, "Report type"), ({"fmt": "pdf"}, "Format"),
                            ({"from_date": "2026-10-05", "to_date": "2026-10-01"}, "end date"),
                            ({"from_date": "last week"}, "YYYY-MM-DD")]:
        with pytest.raises(ReportError, match=message):
            make(service, **kwargs)


def test_role_without_access_gets_a_plain_message(service):
    denied = FakeErp({}, {"reports/sales": ErpError(403, "Your role does not have access to this information.")})
    with pytest.raises(ReportError) as exc:
        make(service, user=WORKER, erp=denied)
    assert exc.value.status == 403 and "role" in exc.value.message


def test_ten_reports_an_hour(service):
    for _ in range(10):
        make(service, kind="receivables")
    with pytest.raises(ReportError) as exc:
        make(service, kind="receivables")
    assert exc.value.status == 429
    make(service, user=SALES, kind="receivables")  # per user
    service.clock[0] = NOW + timedelta(minutes=61)
    make(service, kind="receivables")


def test_download_rules_and_expiry(service):
    report = make(service, user=SALES)
    assert service.file_for(SALES, report["id"])[1] == report["fileName"]
    assert service.file_for(OWNER, report["id"])  # Owner and Admin may download any report
    with pytest.raises(ReportError):
        service.file_for(WORKER, report["id"])
    assert [r["id"] for r in service.list(SALES)] == [report["id"]] and service.list(WORKER) == []
    service.clock[0] = NOW + timedelta(days=7, seconds=1)
    with pytest.raises(ReportError):
        service.file_for(SALES, report["id"])
    assert service.list(SALES) == []


def test_create_report_tool(service):
    output = asyncio.run(run_tool("create_report", {"report_type": "sales", "format": "pptx", "period": "this_month"},
                                  FakeErp(ERP_DATA), ToolContext(OWNER, service)))
    assert output.report["format"] == "pptx" and output.model["fileName"].endswith(".pptx")
    assert "_report" not in output.model and "ready" in output.model["note"]

    denied = FakeErp({}, {"reports/sales": ErpError(403, "Your role does not have access to this information.")})
    output = asyncio.run(run_tool("create_report", {"report_type": "sales"}, denied, ToolContext(WORKER, service)))
    assert "role" in output.model["error"] and output.report is None

    expired = FakeErp({}, {"reports/sales": ErpError(401, "Your session has expired. Sign in again.")})
    with pytest.raises(ErpError):
        asyncio.run(run_tool("create_report", {"report_type": "sales"}, expired, ToolContext(OWNER, service)))


# ---------------------------------------------------------------- through chat and routes

from app import main  # noqa: E402
from app.agent import Agent  # noqa: E402
from app.chat import ChatService  # noqa: E402
from app.erp_client import ResponseCache  # noqa: E402
from app.limits import UserLimits  # noqa: E402
from tests.test_agent import scripted_llm, tool_call  # noqa: E402
from tests.test_chat import KEY, check, client, erp_http, events, token  # noqa: E402,F401


def test_report_asked_for_in_chat_is_linked_to_the_answer(service):
    llm, _ = scripted_llm([tool_call("create_report", {"report_type": "sales", "format": "pptx", "period": "this_month"}),
                           {"role": "assistant", "content": "Your sales report is ready."}])
    store = service._store
    chat = ChatService(store, Agent(llm), llm, UserLimits(20), erp_http(ERP_DATA), ResponseCache(30), service)
    user = check(token())
    session = store.create_session(user.id, user.username)
    got = events(chat.start(user, session["id"], "Make this month's sales as PowerPoint"))

    report = next(d for k, d in got if k == "report")
    assert report["format"] == "pptx" and report["downloadUrl"] == f"/reports/{report['id']}/download"
    assert "userId" not in report
    done = got[-1][1]
    assert done["report"]["id"] == report["id"]
    assert store.messages(session["id"])[1]["reportId"] == report["id"]


def test_report_routes(client):
    main.app.state.erp_http = erp_http(ERP_DATA)
    owner = {"Authorization": f"Bearer {token()}"}
    sales = {"Authorization": f"Bearer {token(2, 'sana', 'Sales')}"}
    worker = {"Authorization": f"Bearer {token(3, 'worker', 'Worker')}"}

    created = client.post("/reports", json={"type": "sales", "format": "xlsx", "from": "2026-10-01", "to": "2026-10-05"},
                          headers=sales)
    assert created.status_code == 201
    report = created.json()
    assert report["fileName"] == "Sales_2026-10-01_to_2026-10-05.xlsx"

    download = client.get(report["downloadUrl"], headers=sales)
    assert download.status_code == 200 and download.content[:2] == b"PK"
    assert "Sales_2026-10-01_to_2026-10-05.xlsx" in download.headers["content-disposition"]
    assert client.get(report["downloadUrl"], headers=owner).status_code == 200
    assert client.get(report["downloadUrl"], headers=worker).status_code == 404
    assert [r["id"] for r in client.get("/reports", headers=sales).json()] == [report["id"]]
    assert client.get("/reports", headers=worker).json() == []
    assert client.post("/reports", json={"type": "payroll"}, headers=sales).status_code == 422
    assert client.get(report["downloadUrl"]).status_code == 401


def test_data_tool_called_with_a_file_format_makes_the_report():
    from app.tools import redirect_file_request
    assert redirect_file_request("sales_report", {"period": "this_month", "format": "xlsx", "report_type": "sales"}) == \
        ("create_report", {"report_type": "sales", "format": "xlsx", "period": "this_month"})
    assert redirect_file_request("production_report", {"format": "PowerPoint"}) == \
        ("create_report", {"report_type": "production", "format": "pptx"})
    assert redirect_file_request("sales_report", {"period": "today"}) == ("sales_report", {"period": "today"})
    assert redirect_file_request("customer_details", {"format": "xlsx"}) == ("customer_details", {"format": "xlsx"})


def test_report_redirect_through_the_agent(service):
    llm, _ = scripted_llm([tool_call("sales_report", {"period": "this_month", "format": "xlsx"}),
                           {"role": "assistant", "content": "Your report is ready."}])
    result = asyncio.run(Agent(llm).ask("Sales report in Excel", [], FakeErp(ERP_DATA), "Owner",
                                        context=ToolContext(OWNER, service)))
    assert [c.name for c in result.tool_calls] == ["create_report"] and result.report["format"] == "xlsx"
