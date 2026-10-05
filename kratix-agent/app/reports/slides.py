"""PowerPoint writer: 16:9, title slide, key figures, native (editable) charts, tables of at most 15 rows per slide."""

from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .catalog import DATE, INT, KG, PERCENT, PKR, QTY, Chart, ReportData, Table, day_text

ROWS_PER_SLIDE = 15
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xEE, 0xF2, 0xF7)
CHART_TYPES = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS}
CHART_NUMBER_FORMATS = {PKR: '#,##0', QTY: '#,##0', KG: '#,##0', INT: '0', PERCENT: '0.0"%"'}


def text_value(value, fmt: str) -> str:
    if value is None or value == "":
        return ""
    if fmt == PKR:
        return f"{value:,.0f}"
    if fmt in (QTY, KG):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    if fmt == PERCENT:
        return f"{value:.1f}%"
    if fmt == INT:
        return f"{value:,.0f}"
    if fmt == DATE:
        return day_text(date.fromisoformat(str(value)[:10]))
    return str(value)


def _text(slide, left, top, width, height, text, size, bold=False, color=None, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    if color is not None:
        paragraph.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align
    return box


def _slide(deck, title: str, company: str):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _text(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8), title, 26, bold=True, color=NAVY)
    _text(slide, Inches(0.5), Inches(7.0), Inches(12.3), Inches(0.4), company, 10, color=MUTED)
    return slide


def _figures(deck, report: ReportData, company: str) -> None:
    slide = _slide(deck, "Key figures", company)
    per_row = 3
    for index, figure in enumerate(report.figures):
        left = Inches(0.5 + (index % per_row) * 4.15)
        top = Inches(1.5 + (index // per_row) * 2.4)
        card = slide.shapes.add_shape(1, left, top, Inches(3.9), Inches(2.1))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.fill.background()
        value = text_value(figure.value, figure.fmt)
        if figure.fmt == PKR:
            value = f"PKR {value}"
        _text(slide, left + Inches(0.2), top + Inches(0.3), Inches(3.5), Inches(1.0), value, 30, bold=True, color=NAVY)
        _text(slide, left + Inches(0.2), top + Inches(1.3), Inches(3.5), Inches(0.6), figure.label, 16, color=MUTED)


def _chart(deck, chart: Chart, company: str) -> None:
    slide = _slide(deck, chart.title, company)
    data = CategoryChartData()
    data.categories = chart.categories
    data.add_series(chart.series, chart.values)
    graphic = slide.shapes.add_chart(CHART_TYPES[chart.kind], Inches(0.7), Inches(1.3), Inches(11.9), Inches(5.6), data).chart
    graphic.has_legend = False
    plot = graphic.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = CHART_NUMBER_FORMATS.get(chart.fmt, "General")
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(11)
    graphic.value_axis.tick_labels.number_format = CHART_NUMBER_FORMATS.get(chart.fmt, "General")
    graphic.value_axis.tick_labels.number_format_is_linked = False
    graphic.value_axis.has_major_gridlines = False
    graphic.category_axis.tick_labels.font.size = Pt(12)
    if chart.kind == "bar":
        graphic.category_axis.reverse_order = True  # largest at the top
    if chart.kind != "line":
        plot.series[0].format.fill.solid()
        plot.series[0].format.fill.fore_color.rgb = NAVY


def _table(deck, table: Table, company: str) -> None:
    chunks = [table.rows[i:i + ROWS_PER_SLIDE] for i in range(0, len(table.rows), ROWS_PER_SLIDE)] or [[]]
    for number, rows in enumerate(chunks, start=1):
        title = table.name if len(chunks) == 1 else f"{table.name} ({number} of {len(chunks)})"
        slide = _slide(deck, title, company)
        if not rows:
            _text(slide, Inches(0.5), Inches(1.5), Inches(12), Inches(0.6), table.empty_text, 18, color=MUTED)
            continue
        shape = slide.shapes.add_table(len(rows) + 1, len(table.columns), Inches(0.5), Inches(1.2), Inches(12.3),
                                       Inches(0.35 * (len(rows) + 1)))
        grid = shape.table
        for col, column in enumerate(table.columns):
            cell = grid.cell(0, col)
            cell.text = column.header
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            font = cell.text_frame.paragraphs[0].font
            font.size, font.bold, font.color.rgb = Pt(11), True, WHITE
        for r, row in enumerate(rows, start=1):
            for col, column in enumerate(table.columns):
                cell = grid.cell(r, col)
                cell.text = text_value(row.get(column.key), column.fmt)
                paragraph = cell.text_frame.paragraphs[0]
                paragraph.font.size = Pt(10)
                if column.fmt in (PKR, QTY, KG, INT, PERCENT):
                    paragraph.alignment = PP_ALIGN.RIGHT


def write_slides(report: ReportData, path: Path, company: str, generated_for: str, generated_on: date) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)

    title = deck.slides.add_slide(deck.slide_layouts[6])
    _text(title, Inches(0.8), Inches(2.2), Inches(11.7), Inches(0.8), company, 24, color=MUTED)
    _text(title, Inches(0.8), Inches(2.9), Inches(11.7), Inches(1.2), report.title, 44, bold=True, color=NAVY)
    _text(title, Inches(0.8), Inches(4.1), Inches(11.7), Inches(0.6), report.period, 22)
    _text(title, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5),
          f"Generated by Kratix Assistant on {day_text(generated_on)} for {generated_for}", 12, color=MUTED)

    _figures(deck, report, company)
    for chart in report.charts:
        _chart(deck, chart, company)
    for table in report.tables:
        _table(deck, table, company)
    deck.save(path)
